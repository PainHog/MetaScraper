"""End-to-end tests driving the CLI over real files on disk."""

import os
import stat
import sys
import wave

from openpyxl import load_workbook

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from metascraper import cli, probe

STUB_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "stub_bin")


def _make_real_wav(path, seconds=1, rate=48000):
    """Write a genuine (tiny) PCM WAV using only the stdlib."""
    with wave.open(path, "wb") as handle:
        handle.setnchannels(1)
        handle.setsampwidth(2)
        handle.setframerate(rate)
        handle.writeframes(b"\x00\x00" * rate * seconds)


def _make_fake_video(path):
    with open(path, "wb") as handle:
        handle.write(b"\x00" * 4096)


def test_cli_fallback_without_tools(tmp_path, monkeypatch):
    """With no ffprobe/exiftool, the CLI still catalogs every file."""
    monkeypatch.setattr(probe.shutil, "which", lambda *_a, **_k: None)

    media = tmp_path / "Footage"
    media.mkdir()
    _make_real_wav(str(media / "take01.wav"))
    _make_fake_video(str(media / "clip.mp4"))

    out = tmp_path / "out"
    code = cli.run([str(media), "-o", str(out), "-q"])
    assert code == 0

    docs = os.listdir(out / "Documents")
    assert sorted(docs) == ["clip.docx", "take01.docx"]
    assert (out / "Master_Catalog.xlsx").exists()
    assert (out / "Master_Catalog.docx").exists()
    assert (out / "Master_Catalog.json").exists()

    workbook = load_workbook(str(out / "Master_Catalog.xlsx"))
    assert workbook["Catalog"].max_row == 3  # header + 2 files


def test_cli_with_stub_ffprobe_extracts_metadata(tmp_path, monkeypatch):
    """A stub ffprobe on PATH exercises the real subprocess path."""
    # Make the stub executable and put it first on PATH.
    stub = os.path.join(STUB_DIR, "ffprobe")
    os.chmod(stub, os.stat(stub).st_mode | stat.S_IEXEC | stat.S_IXGRP | stat.S_IXOTH)
    monkeypatch.setenv("PATH", STUB_DIR + os.pathsep + os.environ.get("PATH", ""))

    media = tmp_path / "Shoot"
    media.mkdir()
    _make_fake_video(str(media / "scene.mp4"))
    _make_real_wav(str(media / "audio.wav"))

    out = tmp_path / "out"
    code = cli.run([str(media), "-o", str(out), "-q"])
    assert code == 0

    workbook = load_workbook(str(out / "Master_Catalog.xlsx"))
    sheet = workbook["Catalog"]
    header = [c.value for c in sheet[1]]
    rows = {}
    for row in sheet.iter_rows(min_row=2, values_only=True):
        record = dict(zip(header, row))
        rows[record["File Name"]] = record

    assert rows["scene.mp4"]["Resolution"] == "1920 × 1080"
    assert rows["scene.mp4"]["Video Codec"] == "h264"
    assert rows["scene.mp4"]["Camera Model"] == "CX-900"
    assert rows["audio.wav"]["Audio Codec"] == "pcm_s16le"
    assert rows["audio.wav"]["Type"] == "Audio"


def test_cli_rerun_updates_not_duplicates(tmp_path, monkeypatch):
    monkeypatch.setattr(probe.shutil, "which", lambda *_a, **_k: None)
    media = tmp_path / "M"
    media.mkdir()
    _make_real_wav(str(media / "a.wav"))
    out = tmp_path / "out"

    cli.run([str(media), "-o", str(out), "-q"])
    # Add a second file and re-run: total should be 2, not 3.
    _make_real_wav(str(media / "b.wav"))
    cli.run([str(media), "-o", str(out), "-q"])

    workbook = load_workbook(str(out / "Master_Catalog.xlsx"))
    assert workbook["Catalog"].max_row == 3  # header + a.wav + b.wav


def test_cli_reports_when_no_media(tmp_path, capsys):
    empty = tmp_path / "empty"
    empty.mkdir()
    code = cli.run([str(empty), "-o", str(tmp_path / "out")])
    assert code == 1
    assert "No media files" in capsys.readouterr().out
