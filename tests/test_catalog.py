"""Tests for the master catalog: upsert semantics and rendering."""

import os
import sys
from datetime import datetime

from openpyxl import load_workbook

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from metascraper.catalog import MasterCatalog
from metascraper.probe import media_info_from_payloads
from tests.sample_payloads import IPHONE_MOV_FFPROBE, WAV_FFPROBE

NOW = datetime(2024, 6, 1, 12, 0, 0)
LATER = datetime(2024, 6, 2, 9, 30, 0)


def _video_info(path="/media/IMG_4021.MOV"):
    return media_info_from_payloads(path, ffprobe_data=IPHONE_MOV_FFPROBE)


def _audio_info(path="/media/ZOOM0001.WAV"):
    return media_info_from_payloads(path, ffprobe_data=WAV_FFPROBE)


def test_upsert_adds_then_updates(tmp_path):
    store = str(tmp_path / "master.json")
    catalog = MasterCatalog(store).load()

    assert catalog.upsert(_video_info(), NOW) is True   # new
    assert catalog.upsert(_audio_info(), NOW) is True    # new
    assert len(catalog.records) == 2

    # Re-cataloging the same path updates rather than duplicates.
    assert catalog.upsert(_video_info(), LATER) is False
    assert len(catalog.records) == 2

    record = catalog.records[os.path.abspath("/media/IMG_4021.MOV")]
    assert record["first_cataloged_at"] == NOW.isoformat()
    assert record["cataloged_at"] == LATER.isoformat()


def test_json_roundtrip_persists_records(tmp_path):
    store = str(tmp_path / "master.json")
    catalog = MasterCatalog(store).load()
    catalog.upsert(_video_info(), NOW)
    catalog.save_json(NOW)

    reloaded = MasterCatalog(store).load()
    assert len(reloaded.records) == 1
    # A brand-new catalog pointed at the same store keeps prior work.
    reloaded.upsert(_audio_info(), LATER)
    assert len(reloaded.records) == 2


def test_summary_totals(tmp_path):
    catalog = MasterCatalog(str(tmp_path / "m.json")).load()
    catalog.upsert(_video_info(), NOW)
    catalog.upsert(_audio_info(), NOW)
    summary = catalog.summary()
    assert summary["count"] == 2
    assert summary["by_kind"]["Video"] == 1
    assert summary["by_kind"]["Audio"] == 1
    assert summary["total_bytes"] == 74691234 + 1008000


def test_render_xlsx_is_readable(tmp_path):
    catalog = MasterCatalog(str(tmp_path / "m.json")).load()
    catalog.upsert(_video_info(), NOW)
    catalog.upsert(_audio_info(), NOW)
    xlsx = str(tmp_path / "Master.xlsx")
    catalog.render_xlsx(xlsx, NOW)

    workbook = load_workbook(xlsx)
    assert "Catalog" in workbook.sheetnames
    assert "Summary" in workbook.sheetnames
    sheet = workbook["Catalog"]
    header = [c.value for c in sheet[1]]
    assert "File Name" in header and "Resolution" in header
    # header + two data rows
    assert sheet.max_row == 3
    assert sheet.freeze_panes == "A2"


def test_render_docx_is_written(tmp_path):
    from docx import Document

    catalog = MasterCatalog(str(tmp_path / "m.json")).load()
    catalog.upsert(_video_info(), NOW)
    docx_path = str(tmp_path / "Master.docx")
    catalog.render_docx(docx_path, NOW)
    assert os.path.getsize(docx_path) > 0
    document = Document(docx_path)
    text = "\n".join(p.text for p in document.paragraphs)
    assert "Master Media Catalog" in text


def test_corrupt_store_is_backed_up_not_fatal(tmp_path):
    store = tmp_path / "master.json"
    store.write_text("{ this is not valid json ")
    catalog = MasterCatalog(str(store)).load()
    assert catalog.records == {}
    assert (tmp_path / "master.json.corrupt").exists()
