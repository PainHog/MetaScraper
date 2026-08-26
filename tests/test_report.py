"""Tests for the per-file Word report generator."""

import os
import sys
from datetime import datetime

from docx import Document

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from metascraper.probe import media_info_from_payloads
from metascraper.report import write_report
from tests.sample_payloads import IPHONE_EXIFTOOL, IPHONE_MOV_FFPROBE

NOW = datetime(2024, 6, 1, 12, 0, 0)


def _all_text(document) -> str:
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)
    return "\n".join(parts)


def test_report_written_and_contains_key_facts(tmp_path):
    info = media_info_from_payloads(
        "IMG_4021.MOV",
        ffprobe_data=IPHONE_MOV_FFPROBE,
        exiftool_data=IPHONE_EXIFTOOL,
    )
    out = str(tmp_path / "report.docx")
    write_report(info, out, generated_at=NOW)
    assert os.path.getsize(out) > 0

    text = _all_text(Document(out))
    assert "IMG_4021.MOV" in text
    assert "Project Information" in text
    assert "h264" in text or "H.264" in text
    assert "3840 × 2160" in text
    assert "iPhone 15 Pro" in text
    # The full-metadata appendix should preserve raw keys.
    assert "Complete Technical Metadata" in text
    assert "major_brand" in text


def test_report_handles_missing_metadata(tmp_path):
    info = media_info_from_payloads(
        "silent.mov",
        ffprobe_data={},
        notes=["ffprobe was not found on this system"],
    )
    out = str(tmp_path / "basic.docx")
    write_report(info, out, generated_at=NOW)
    text = _all_text(Document(out))
    assert "silent.mov" in text
    assert "ffprobe was not found" in text
