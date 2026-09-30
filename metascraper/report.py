"""Generate a polished per-file "Project Information" Word document."""

from __future__ import annotations

import os
from datetime import datetime
from typing import Any, Dict, List, Tuple

from . import docx_style as ds
from . import utils
from .models import MediaInfo, StreamSummary


def _at_a_glance_rows(info: MediaInfo) -> List[Tuple[str, str]]:
    video = info.primary_video
    audio = info.primary_audio
    rows: List[Tuple[str, str]] = [
        ("Media Type", info.media_kind),
        ("Duration", info.duration_human),
        ("File Size", info.size_human),
        ("Container", utils.coalesce(info.container_long_name, info.container_format)),
    ]
    if video:
        rows.append(("Resolution", utils.coalesce(video.resolution)))
        rows.append(("Frame Rate", utils.human_frame_rate(video.frame_rate)))
    if audio:
        rows.append(("Audio", _audio_glance(audio)))
    if info.recorded_at:
        rows.append(("Recorded", utils.format_datetime(info.recorded_at)))
    return rows


def _audio_glance(audio: StreamSummary) -> str:
    parts = []
    if audio.codec_name:
        parts.append(audio.codec_name.upper())
    if audio.sample_rate:
        parts.append(utils.human_sample_rate(audio.sample_rate))
    if audio.channels:
        parts.append(_channel_label(audio.channels, audio.channel_layout))
    return " · ".join(parts) if parts else "—"


def _channel_label(channels: int, layout: str = None) -> str:
    if layout and layout not in ("", "unknown"):
        return layout
    return {1: "Mono", 2: "Stereo"}.get(channels, f"{channels} ch")


def _file_rows(info: MediaInfo) -> List[Tuple[str, str]]:
    return [
        ("File Name", info.name),
        ("Folder", os.path.basename(os.path.dirname(info.path)) or "—"),
        ("Full Path", info.path),
        ("Media Type", info.media_kind),
        ("File Size", f"{info.size_human}"
            + (f"  ({info.size_bytes:,} bytes)" if info.size_bytes else "")),
        ("Date Created", utils.format_datetime(info.fs_created)),
        ("Date Modified", utils.format_datetime(info.fs_modified)),
    ]


def _container_rows(info: MediaInfo) -> List[Tuple[str, str]]:
    rows = [
        ("Format", utils.coalesce(info.container_long_name, info.container_format)),
        ("Format Name", utils.coalesce(info.container_format)),
        ("Duration", f"{info.duration_human}"
            + (f"  ({info.duration_seconds:.3f} s)" if info.duration_seconds else "")),
        ("Overall Bitrate", utils.human_bitrate(info.overall_bit_rate)),
        ("Streams", str(len(info.streams)) if info.streams else "—"),
    ]
    if info.title:
        rows.append(("Title", info.title))
    if info.artist:
        rows.append(("Artist / Author", info.artist))
    if info.comment:
        rows.append(("Comment", info.comment))
    return rows


def _video_rows(video: StreamSummary) -> List[Tuple[str, str]]:
    rows = [
        ("Codec", utils.coalesce(video.codec_long_name, video.codec_name)),
        ("Profile", utils.coalesce(video.profile)),
        ("Resolution", utils.coalesce(video.resolution)),
        ("Aspect Ratio", utils.coalesce(video.display_aspect_ratio)),
        ("Frame Rate", utils.human_frame_rate(video.frame_rate)),
        ("Bitrate", utils.human_bitrate(video.bit_rate)),
        ("Pixel Format", utils.coalesce(video.pixel_format)),
    ]
    if video.color_space or video.color_primaries or video.color_transfer:
        color = " / ".join(
            x for x in (video.color_space, video.color_primaries, video.color_transfer)
            if x
        )
        rows.append(("Color", color))
    if video.field_order:
        rows.append(("Scan Type", video.field_order))
    if video.rotation:
        rows.append(("Rotation", f"{video.rotation}°"))
    if video.handler_name:
        rows.append(("Handler", video.handler_name))
    return rows


def _audio_rows(audio: StreamSummary) -> List[Tuple[str, str]]:
    rows = [
        ("Codec", utils.coalesce(audio.codec_long_name, audio.codec_name)),
        ("Profile", utils.coalesce(audio.profile)),
        ("Sample Rate", utils.human_sample_rate(audio.sample_rate)),
        ("Channels", _channel_label(audio.channels, audio.channel_layout)
            if audio.channels else "—"),
        ("Bitrate", utils.human_bitrate(audio.bit_rate)),
    ]
    if audio.bits_per_sample:
        rows.append(("Bit Depth", f"{audio.bits_per_sample}-bit"))
    if audio.language:
        rows.append(("Language", audio.language))
    if audio.handler_name:
        rows.append(("Handler", audio.handler_name))
    return rows


def _device_rows(info: MediaInfo) -> List[Tuple[str, str]]:
    return [
        ("Camera / Device Make", utils.coalesce(info.camera_make)),
        ("Camera / Device Model", utils.coalesce(info.camera_model)),
        ("Lens", utils.coalesce(info.lens_model)),
        ("Software", utils.coalesce(info.software)),
        ("Recorded", utils.format_datetime(info.recorded_at)
            if info.recorded_at else "—"),
        ("GPS Location", utils.coalesce(info.gps_coordinates)),
    ]


def _has_device_info(info: MediaInfo) -> bool:
    return any([
        info.camera_make, info.camera_model, info.lens_model,
        info.software, info.recorded_at, info.gps_coordinates,
    ])


def _flatten(data: Any, prefix: str = "") -> List[Tuple[str, str]]:
    """Flatten nested dict/list structures into ``(dotted-key, value)`` rows."""
    rows: List[Tuple[str, str]] = []
    if isinstance(data, dict):
        for key, value in data.items():
            new_prefix = f"{prefix}.{key}" if prefix else str(key)
            rows.extend(_flatten(value, new_prefix))
    elif isinstance(data, list):
        for idx, value in enumerate(data):
            new_prefix = f"{prefix}[{idx}]"
            rows.extend(_flatten(value, new_prefix))
    else:
        rows.append((prefix, "" if data is None else str(data)))
    return rows


def _append_full_metadata(document, info: MediaInfo) -> None:
    """Dump every raw key/value so nothing is lost from the catalog."""
    ds.add_section_heading(document, "Complete Technical Metadata")
    ds.add_caption(
        document,
        "Every field reported by the metadata tools, preserved verbatim for archival.",
    )

    fmt = info.ffprobe_raw.get("format", {}) or {}
    fmt_rows = [(k, v) for k, v in _flatten(fmt) if k != "filename"]
    if fmt_rows:
        _appendix_subheading(document, "Container / Format")
        ds.add_data_table(document, ["Field", "Value"], fmt_rows,
                          col_widths_in=[2.6, 4.0])

    for stream in info.ffprobe_raw.get("streams", []) or []:
        idx = stream.get("index", "?")
        codec_type = stream.get("codec_type", "stream").title()
        _appendix_subheading(document, f"Stream #{idx} — {codec_type}")
        rows = _flatten(stream)
        ds.add_data_table(document, ["Field", "Value"], rows,
                          col_widths_in=[2.6, 4.0])

    if info.exiftool_raw:
        exif_rows = [
            (k, v) for k, v in _flatten(info.exiftool_raw)
            if k not in ("SourceFile",)
        ]
        if exif_rows:
            _appendix_subheading(document, "ExifTool")
            ds.add_data_table(document, ["Field", "Value"], exif_rows,
                              col_widths_in=[2.6, 4.0])


def _appendix_subheading(document, text: str) -> None:
    from docx.shared import Pt

    p = document.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(3)
    run = p.add_run(utils.xml_safe(text))
    run.font.bold = True
    run.font.size = Pt(10)
    run.font.color.rgb = ds.BRAND


def build_report_document(info: MediaInfo, generated_at: datetime):
    """Build and return a python-docx Document for a single media file."""
    document = ds.new_document()

    ds.add_title_block(
        document,
        title=info.name,
        subtitle="Project Information — Media Metadata Catalog",
        eyebrow="Media Asset Record",
    )

    # Generated-on line.
    from docx.shared import Pt

    meta_p = document.add_paragraph()
    meta_run = meta_p.add_run(
        f"Cataloged {utils.format_datetime(generated_at)}"
    )
    meta_run.font.size = Pt(9)
    meta_run.font.color.rgb = ds.MUTED
    meta_p.paragraph_format.space_after = Pt(6)

    # At a glance.
    ds.add_section_heading(document, "At a Glance")
    ds.add_key_value_table(document, _at_a_glance_rows(info))

    # Notes / warnings.
    if info.notes:
        ds.add_section_heading(document, "Notes")
        for note in info.notes:
            p = document.add_paragraph(style=None)
            run = p.add_run(utils.xml_safe(f"• {note}"))
            run.font.size = Pt(9.5)
            run.font.color.rgb = ds.ACCENT

    # File details.
    ds.add_section_heading(document, "File Details")
    ds.add_key_value_table(document, _file_rows(info))

    # Container.
    if info.container_format or info.duration_seconds:
        ds.add_section_heading(document, "Container / Format")
        ds.add_key_value_table(document, _container_rows(info))

    # Video streams.
    for i, video in enumerate(info.video_streams):
        suffix = f" #{i + 1}" if len(info.video_streams) > 1 else ""
        ds.add_section_heading(document, f"Video Stream{suffix}")
        ds.add_key_value_table(document, _video_rows(video))

    # Audio streams.
    for i, audio in enumerate(info.audio_streams):
        suffix = f" #{i + 1}" if len(info.audio_streams) > 1 else ""
        ds.add_section_heading(document, f"Audio Stream{suffix}")
        ds.add_key_value_table(document, _audio_rows(audio))

    # Device / provenance.
    if _has_device_info(info):
        ds.add_section_heading(document, "Camera / Device & Provenance")
        ds.add_key_value_table(document, _device_rows(info))

    # Full raw dump.
    if info.ffprobe_raw or info.exiftool_raw:
        _append_full_metadata(document, info)

    ds.add_footer(document, "MetaScraper · Media Metadata Catalog")
    return document


def write_report(info: MediaInfo, output_path: str, generated_at: datetime = None) -> str:
    """Build and save a per-file report, returning the output path."""
    generated_at = generated_at or datetime.now().astimezone()
    document = build_report_document(info, generated_at)
    os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)
    document.save(output_path)
    return output_path
